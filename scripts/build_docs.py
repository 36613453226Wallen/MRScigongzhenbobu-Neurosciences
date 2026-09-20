#!/usr/bin/env python3
"""Build formatted DOCX and HTML (for Chrome PDF) from the research notes."""
from __future__ import annotations

from pathlib import Path
import html
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path("/workspace")
DOCS = ROOT / "docs"
OUT = ROOT / "docs" / "export"
OUT.mkdir(parents=True, exist_ok=True)

FILES = {
    "src": DOCS / "00_文献来源标注.md",
    "qa": DOCS / "01_问答详解_论文规格结论与性别差异.md",
    "adapt": DOCS / "02_长期结构改变_行事思考与适应方向.md",
    "struct": DOCS / "03_长期结构改变_文献笔记.md",
    "testo": DOCS / "04_睾酮保护原理_文献笔记.md",
}

INLINE_RE = re.compile(r"(\*\*[^*]+?\*\*|\[[^\]]+\]\([^)]+\))")
LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def set_run_font(run, name="Microsoft YaHei", size=11, bold=False, color=None, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = name
    if color:
        run.font.color.rgb = color
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    rFonts.set(qn("w:eastAsia"), "WenQuanYi Micro Hei")
    rFonts.set(qn("w:cs"), "Calibri")


def configure_styles(doc: Document):
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "WenQuanYi Micro Hei")
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(8)


def add_heading_para(doc, text, level):
    p = doc.add_paragraph()
    if level == 0:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, size=18, bold=True, color=RGBColor(0x1F, 0x3A, 0x5F))
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.line_spacing = 1.3
    elif level == 1:
        run = p.add_run(text)
        set_run_font(run, size=16, bold=True, color=RGBColor(0x1F, 0x3A, 0x5F))
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
    elif level == 2:
        run = p.add_run(text)
        set_run_font(run, size=13.5, bold=True, color=RGBColor(0x2C, 0x5F, 0x7C))
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
    else:
        run = p.add_run(text)
        set_run_font(run, size=12, bold=True, color=RGBColor(0x3D, 0x5A, 0x40))
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
    return p


def add_hyperlink(paragraph, text, url, size=11, bold=False):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1A5F9E")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    if bold:
        rPr.append(OxmlElement("w:b"))
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rPr.append(sz)
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    rFonts.set(qn("w:eastAsia"), "WenQuanYi Micro Hei")
    rPr.append(rFonts)
    new_run.append(rPr)
    text_elem = OxmlElement("w:t")
    text_elem.set(qn("xml:space"), "preserve")
    text_elem.text = text
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def fill_md_runs(paragraph, text, size=11, default_bold=False, color=None):
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos:m.start()])
            set_run_font(run, size=size, bold=default_bold, color=color or RGBColor(0x22, 0x22, 0x22))
        token = m.group(0)
        link = LINK_RE.fullmatch(token)
        if link:
            add_hyperlink(paragraph, link.group(1), link.group(2), size=size, bold=default_bold)
        else:
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True, color=color or RGBColor(0x22, 0x22, 0x22))
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=size, bold=default_bold, color=color or RGBColor(0x22, 0x22, 0x22))


def add_body(doc, text, *, bold=False, italic=False, quote=False, bullet=False):
    p = doc.add_paragraph()
    if bullet:
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        text = "•  " + text
    if quote:
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.right_indent = Cm(0.5)
    fill_md_runs(p, text, size=11, default_bold=bold, color=RGBColor(0x33, 0x33, 0x33) if italic else RGBColor(0x22, 0x22, 0x22))
    if italic:
        for run in p.runs:
            run.italic = True
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    return p


def add_table_from_md(doc, rows):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.rows[i].cells[j]
            val = row[j] if j < len(row) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            fill_md_runs(
                p,
                val,
                size=9,
                default_bold=(i == 0),
                color=RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else RGBColor(0x22, 0x22, 0x22),
            )
            if i == 0:
                shading = cell._element.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "1F3A5F")
                shd.set(qn("w:val"), "clear")
                shading.append(shd)
    doc.add_paragraph()


def md_to_docx_parts(doc, md_text: str, skip_first_h1=False):
    lines = md_text.splitlines()
    i = 0
    skipped_h1 = False
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.strip() == "---":
            i += 1
            continue
        if line.startswith("|") and "|" in line[1:]:
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in raw):
                    rows.append(raw)
                i += 1
            add_table_from_md(doc, rows)
            continue
        if line.startswith("# "):
            if skip_first_h1 and not skipped_h1:
                skipped_h1 = True
                add_heading_para(doc, line[2:].strip(), 0)
            else:
                add_heading_para(doc, line[2:].strip(), 1)
            i += 1
            continue
        if line.startswith("## "):
            add_heading_para(doc, line[3:].strip(), 1 if skip_first_h1 else 2)
            i += 1
            continue
        if line.startswith("### "):
            add_heading_para(doc, line[4:].strip(), 2 if skip_first_h1 else 3)
            i += 1
            continue
        if line.startswith("> "):
            buf = []
            while i < len(lines) and lines[i].startswith("> "):
                buf.append(lines[i][2:].strip())
                i += 1
            add_body(doc, " ".join(buf), italic=True, quote=True)
            continue
        if re.match(r"^\d+\.\s+", line) or line.startswith("- "):
            text = re.sub(r"^\d+\.\s+", "", line)
            text = re.sub(r"^- ", "", text)
            add_body(doc, text, bullet=True)
            i += 1
            continue
        buf = [line]
        i += 1
        while (
            i < len(lines)
            and lines[i].strip()
            and not lines[i].startswith(("#", "|", "-", ">", "---"))
            and not re.match(r"^\d+\.\s+", lines[i])
        ):
            buf.append(lines[i].strip())
            i += 1
        para = " ".join(buf)
        add_body(doc, para)


def cover_page(doc, title, subtitle, extra_lines):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("文献收录")
    set_run_font(run, size=14, color=RGBColor(0x6B, 0x7C, 0x8A))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    set_run_font(run, size=22, bold=True, color=RGBColor(0x1F, 0x3A, 0x5F))
    p.paragraph_format.space_after = Pt(16)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle)
    set_run_font(run, size=13, color=RGBColor(0x2C, 0x5F, 0x7C))
    for line in extra_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        set_run_font(run, size=11, color=RGBColor(0x44, 0x44, 0x44))
    doc.add_page_break()


def write_docx(path: Path, title, subtitle, extras, parts):
    doc = Document()
    configure_styles(doc)
    cover_page(doc, title, subtitle, extras)
    for idx, (md_path, label) in enumerate(parts):
        text = md_path.read_text(encoding="utf-8")
        md_to_docx_parts(doc, text, skip_first_h1=True)
        if idx != len(parts) - 1:
            doc.add_page_break()
    doc.save(path)
    print("wrote", path)


# ---------- HTML / PDF ----------

CSS = """
@page { size: A4; margin: 2.1cm 2.0cm 2.2cm 2.0cm; }
html, body { font-family: "WenQuanYi Micro Hei", "Noto Sans CJK SC", "Droid Sans Fallback", sans-serif;
  color: #222; font-size: 11.5pt; line-height: 1.62; }
body { max-width: 100%; }
h1 { font-size: 20pt; color: #1F3A5F; line-height: 1.35; margin: 0.2em 0 0.6em; }
h2 { font-size: 15pt; color: #1F3A5F; border-bottom: 1.5px solid #d5dde6; padding-bottom: 0.18em; margin-top: 1.4em; }
h3 { font-size: 13pt; color: #2C5F7C; margin-top: 1.1em; }
h4 { font-size: 12pt; color: #3D5A40; }
p { margin: 0.35em 0 0.7em; }
blockquote { margin: 0.8em 1.2em; padding: 0.4em 0.9em; border-left: 4px solid #2C5F7C;
  color: #333; font-style: italic; background: #f4f7fa; }
ul, ol { margin: 0.3em 0 0.8em 1.4em; }
li { margin: 0.22em 0; }
table { border-collapse: collapse; width: 100%; margin: 0.6em 0 1.1em; font-size: 10pt; }
th { background: #1F3A5F; color: #fff; text-align: left; padding: 6px 8px; }
td { border: 1px solid #cfd8e3; padding: 6px 8px; vertical-align: top; }
tr:nth-child(even) td { background: #f6f8fb; }
hr { border: 0; border-top: 1px solid #d5dde6; margin: 1.2em 0; }
.cover { page-break-after: always; padding-top: 4.2cm; text-align: center; }
.cover .kicker { letter-spacing: 0.4em; color: #6B7C8A; font-size: 12pt; margin-bottom: 1.2em; }
.cover h1 { font-size: 26pt; margin: 0.4em 1em 0.6em; }
.cover .sub { color: #2C5F7C; font-size: 13.5pt; margin: 0.4em 1.5em 1.4em; }
.cover .meta { color: #555; font-size: 11pt; line-height: 1.8; }
.section-break { page-break-before: always; }
code { font-family: "WenQuanYi Micro Hei Mono", monospace; font-size: 0.92em; }
a { color: #125a9e; text-decoration: underline; }
"""


def _fmt_plain(s: str) -> str:
    s = html.escape(s)
    s = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", s)
    s = re.sub(r"`(.+?)`", r"<code>\1</code>", s)
    return s


def md_inline(s: str) -> str:
    out = []
    last = 0
    for m in LINK_RE.finditer(s):
        out.append(_fmt_plain(s[last:m.start()]))
        href = html.escape(m.group(2), quote=True)
        out.append(f'<a href="{href}">{_fmt_plain(m.group(1))}</a>')
        last = m.end()
    out.append(_fmt_plain(s[last:]))
    return "".join(out)


def md_to_html_body(md_text: str) -> str:
    lines = md_text.splitlines()
    out = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.strip() == "---":
            out.append("<hr/>")
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in raw):
                    rows.append(raw)
                i += 1
            if rows:
                head, body = rows[0], rows[1:]
                out.append("<table><thead><tr>" + "".join(f"<th>{md_inline(c)}</th>" for c in head) + "</tr></thead><tbody>")
                for r in body:
                    out.append("<tr>" + "".join(f"<td>{md_inline(c)}</td>" for c in r) + "</tr>")
                out.append("</tbody></table>")
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            lv = len(m.group(1))
            out.append(f"<h{lv}>{md_inline(m.group(2))}</h{lv}>")
            i += 1
            continue
        if line.startswith("> "):
            buf = []
            while i < len(lines) and lines[i].startswith("> "):
                buf.append(lines[i][2:].strip())
                i += 1
            out.append("<blockquote>" + md_inline(" ".join(buf)) + "</blockquote>")
            continue
        if line.startswith("- "):
            out.append("<ul>")
            while i < len(lines) and lines[i].startswith("- "):
                out.append("<li>" + md_inline(lines[i][2:].strip()) + "</li>")
                i += 1
            out.append("</ul>")
            continue
        if re.match(r"^\d+\.\s+", line):
            out.append("<ol>")
            while i < len(lines) and re.match(r"^\d+\.\s+", lines[i]):
                out.append("<li>" + md_inline(re.sub(r"^\d+\.\s+", "", lines[i]).strip()) + "</li>")
                i += 1
            out.append("</ol>")
            continue
        buf = [line.strip()]
        i += 1
        while (
            i < len(lines)
            and lines[i].strip()
            and not lines[i].startswith(("#", "|", "-", ">", "---"))
            and not re.match(r"^\d+\.\s+", lines[i])
            and lines[i].strip() != "---"
        ):
            buf.append(lines[i].strip())
            i += 1
        out.append("<p>" + md_inline(" ".join(buf)) + "</p>")
    return "\n".join(out)


def html_doc(title, subtitle, extras, sections):
    cover = f"""
    <section class="cover">
      <div class="kicker">文献收录</div>
      <h1>{html.escape(title)}</h1>
      <div class="sub">{html.escape(subtitle)}</div>
      <div class="meta">{"<br/>".join(html.escape(x) for x in extras)}</div>
    </section>
    """
    bodies = []
    for idx, (path, _label) in enumerate(sections):
        cls = "section-break" if idx else ""
        bodies.append(f'<section class="{cls}">' + md_to_html_body(path.read_text(encoding="utf-8")) + "</section>")
    return f"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8"/>
<title>{html.escape(title)}</title>
<style>{CSS}</style></head><body>
{cover}
{''.join(bodies)}
</body></html>"""


def main():
    common_extra = [
        "对应讨论原文：Savic, Cerebral Cortex, 2020, 30(6):3759–3770",
        "DOI: 10.1093/cercor/bhz340",
        "收录日期：2026-09-20",
        "体例：出处 / 简短数据 / 四点体现分析 / 主要结论",
    ]

    write_docx(
        OUT / "慢性应激大脑改变与睾酮保护_完整收录.docx",
        "慢性应激的长期结构改变与睾酮保护",
        "问答详解 · 适应方向 · 境外与境内文献笔记",
        common_extra,
        [
            (FILES["src"], "来源"),
            (FILES["qa"], "问答"),
            (FILES["adapt"], "适应"),
            (FILES["struct"], "结构"),
            (FILES["testo"], "睾酮"),
        ],
    )
    write_docx(
        OUT / "长期结构改变_文献笔记.docx",
        "长期结构改变",
        "境外 4 篇 + 大陆境内 2 篇 文献笔记",
        common_extra + ["主题文档 3.1"],
        [(FILES["src"], "来源"), (FILES["struct"], "结构")],
    )
    write_docx(
        OUT / "睾酮保护原理_文献笔记.docx",
        "睾酮保护原理",
        "境外 4 篇 + 大陆境内 2 篇 文献笔记",
        common_extra + ["主题文档 3.2"],
        [(FILES["src"], "来源"), (FILES["testo"], "睾酮")],
    )

    html_specs = [
        (
            OUT / "慢性应激大脑改变与睾酮保护_完整收录.html",
            "慢性应激的长期结构改变与睾酮保护",
            "问答详解 · 适应方向 · 境外与境内文献笔记",
            common_extra,
            [
                (FILES["src"], "来源"),
                (FILES["qa"], "问答"),
                (FILES["adapt"], "适应"),
                (FILES["struct"], "结构"),
                (FILES["testo"], "睾酮"),
            ],
        ),
        (
            OUT / "长期结构改变_文献笔记.html",
            "长期结构改变",
            "境外 4 篇 + 大陆境内 2 篇 文献笔记",
            common_extra + ["主题文档 3.1"],
            [(FILES["src"], "来源"), (FILES["struct"], "结构")],
        ),
        (
            OUT / "睾酮保护原理_文献笔记.html",
            "睾酮保护原理",
            "境外 4 篇 + 大陆境内 2 篇 文献笔记",
            common_extra + ["主题文档 3.2"],
            [(FILES["src"], "来源"), (FILES["testo"], "睾酮")],
        ),
    ]
    for path, title, subtitle, extras, sections in html_specs:
        path.write_text(html_doc(title, subtitle, extras, sections), encoding="utf-8")
        print("wrote", path)

    archive = ROOT / "收录"
    archive.mkdir(parents=True, exist_ok=True)
    for name in (
        "慢性应激大脑改变与睾酮保护_完整收录.docx",
        "长期结构改变_文献笔记.docx",
        "睾酮保护原理_文献笔记.docx",
    ):
        target = archive / name
        target.write_bytes((OUT / name).read_bytes())
        print("copied", target)


if __name__ == "__main__":
    main()
